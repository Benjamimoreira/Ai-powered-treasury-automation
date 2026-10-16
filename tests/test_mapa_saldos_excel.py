from datetime import date

import openpyxl

from app.db.models import SaldoDiario
from app.services import mapa_saldos_excel
from app.services.mapa_saldos_excel import contas_mapa_saldos, saldo_total, serie_saldo_total_bruta, totais_mapa_saldos


def _criar_mapa(raiz, ano, mes, folhas):
    """folhas: {"DD-MM": [(banco, entidade, contabilistico, disponivel), ...]}
    - mesmo formato do Mapa de Saldos real (cabeçalho na linha 2, contas a
    partir da 3, linha "Saldo" com =SUM no fim)."""
    pasta = raiz / "CONTABILIDADE" / "08- Saldos Bancários" / str(ano)
    pasta.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome, contas in folhas.items():
        ws = wb.create_sheet(nome)
        ws["A1"] = f"SALDOS BANCÁRIOS {nome}"
        ws.append(["Nº Ordem", "Banco", "Entidade", "NIB", "Saldo Contabilistico", "Disponível", "Observações"])
        for banco, entidade, contabilistico, disponivel in contas:
            ws.append([None, banco, entidade, "0035", contabilistico, disponivel])
        fim = ws.max_row
        ws.append([])
        ws.append([None, None, None, "Saldo", f"=SUM(E3:E{fim})", f"=SUM(F3:F{fim})"])
        ws.append([None, None, None, "Total", 999999, 999999])  # não conta
    wb.save(pasta / f"{mes:02d}_Mapa de Saldos de mes.xlsx")
    mapa_saldos_excel._cache.clear()


def test_soma_as_contas_da_folha_do_dia_incluindo_outros_bancos(tmp_path):
    _criar_mapa(tmp_path, 2026, 10, {
        "16-10": [("CGD", "HCN", 1000.5, 900.0), ("Santander ", "HII", 5000, 5000), ("CGD", "F&C", "n/d", None)],
        "15-10": [("CGD", "HCN", 1.0, 1.0)],
    })

    totais = totais_mapa_saldos(date(2026, 10, 16), raiz=str(tmp_path))

    assert totais == {
        "entidades": 3,
        "saldo_contabilistico_total": 6000.5,
        "saldo_disponivel_total": 5900.0,
        "fonte": "mapa_saldos",
        "dia_referencia": "2026-10-16",
    }


def test_sem_folha_do_dia_usa_a_ultima_anterior(tmp_path):
    _criar_mapa(tmp_path, 2026, 10, {"09-10": [("CGD", "HCN", 10.0, 10.0)], "13-10": [("CGD", "HCN", 20.0, 20.0)]})

    totais = totais_mapa_saldos(date(2026, 10, 12), raiz=str(tmp_path))

    assert totais["dia_referencia"] == "2026-10-09"
    assert totais["saldo_contabilistico_total"] == 10.0


def test_inicio_do_mes_sem_folhas_usa_o_mapa_do_mes_anterior(tmp_path):
    _criar_mapa(tmp_path, 2026, 9, {"30-09": [("CGD", "HCN", 30.0, 30.0)]})
    _criar_mapa(tmp_path, 2026, 10, {"02-10": [("CGD", "HCN", 40.0, 40.0)]})

    totais = totais_mapa_saldos(date(2026, 10, 1), raiz=str(tmp_path))

    assert totais["dia_referencia"] == "2026-09-30"


def test_saldo_total_sem_mapa_usa_os_extratos(db_session, tmp_path, monkeypatch):
    monkeypatch.setenv("ONEDRIVE_RAIZ", str(tmp_path))
    db_session.add(SaldoDiario(dia=date(2026, 10, 16), entidade="A", saldo_contabilistico=7.0, saldo_disponivel=6.0))
    db_session.commit()

    resultado = saldo_total(db_session, date(2026, 10, 16))

    assert resultado["fonte"] == "extratos"
    assert resultado["saldo_contabilistico_total"] == 7.0


def test_saldo_total_prefere_o_mapa(db_session, tmp_path, monkeypatch):
    monkeypatch.setenv("ONEDRIVE_RAIZ", str(tmp_path))
    _criar_mapa(tmp_path, 2026, 10, {"16-10": [("CGD", "HCN", 100.0, 90.0)]})
    db_session.add(SaldoDiario(dia=date(2026, 10, 16), entidade="Conta fechada", saldo_contabilistico=7.0, saldo_disponivel=6.0))
    db_session.commit()

    resultado = saldo_total(db_session, date(2026, 10, 16))

    assert resultado["fonte"] == "mapa_saldos"
    assert resultado["saldo_contabilistico_total"] == 100.0


def test_endpoint_saldo_total_devolve_fonte(client, tmp_path, monkeypatch):
    monkeypatch.setenv("ONEDRIVE_RAIZ", str(tmp_path))
    _criar_mapa(tmp_path, 2026, 10, {"16-10": [("CGD", "HCN", 100.0, 90.0)]})

    r = client.get("/saldo-total", params={"dia": "2026-10-16"})

    assert r.status_code == 200
    assert r.json()["fonte"] == "mapa_saldos"
    assert r.json()["dia_referencia"] == "2026-10-16"
    assert r.json()["saldo_contabilistico_total"] == 100.0


def test_contas_com_variacao_face_a_folha_anterior(tmp_path):
    _criar_mapa(tmp_path, 2026, 10, {
        "16-10": [("CGD", "HCN", 150.0, 150.0), ("Abanca -", "HII", 10000, 10000), ("CGD", "Nova", 5.0, 5.0)],
        "15-10": [("CGD", "HCN", 100.0, 90.0), ("Abanca -", "HII", 10000, 10000)],
    })

    resultado = contas_mapa_saldos(date(2026, 10, 16), raiz=str(tmp_path))

    assert resultado["dia_referencia"] == "2026-10-16"
    assert resultado["dia_anterior"] == "2026-10-15"
    assert resultado["saldo_contabilistico_total"] == 10155.0
    hcn, hii, nova = resultado["contas"]
    assert (hcn["banco"], hcn["entidade"]) == ("CGD", "HCN")
    assert hcn["variacao_contabilistico"] == 50.0
    assert hcn["variacao_pct_contabilistico"] == 50.0
    assert hcn["variacao_disponivel"] == 60.0
    assert hii["variacao_contabilistico"] == 0
    assert nova["variacao_contabilistico"] is None  # conta nova, sem folha anterior


def test_serie_usa_o_mapa_e_os_extratos_antes_de_haver_mapa(db_session, tmp_path):
    db_session.add(SaldoDiario(dia=date(2026, 9, 1), entidade="A", saldo_contabilistico=1.0, saldo_disponivel=1.0))
    db_session.add(SaldoDiario(dia=date(2026, 10, 16), entidade="A", saldo_contabilistico=2.0, saldo_disponivel=2.0))
    db_session.commit()
    _criar_mapa(tmp_path, 2026, 10, {"15-10": [("CGD", "HCN", 100.0, 90.0)], "16-10": [("CGD", "HCN", 200.0, 190.0)]})

    serie = serie_saldo_total_bruta(db_session, raiz=str(tmp_path))

    assert [(p["dia"], p["saldo_contabilistico_total"]) for p in serie] == [
        (date(2026, 9, 1), 1.0),     # ainda sem Mapa: extratos
        (date(2026, 10, 15), 100.0),  # folha do Mapa
        (date(2026, 10, 16), 200.0),  # folha do Mapa, não os 2.0 dos extratos
    ]


def test_endpoint_mapa_saldos_sem_mapa_devolve_404(client, tmp_path, monkeypatch):
    monkeypatch.setenv("ONEDRIVE_RAIZ", str(tmp_path))

    assert client.get("/saldos/mapa-saldos", params={"dia": "2026-10-16"}).status_code == 404


def test_endpoint_mapa_saldos_devolve_as_contas(client, tmp_path, monkeypatch):
    monkeypatch.setenv("ONEDRIVE_RAIZ", str(tmp_path))
    _criar_mapa(tmp_path, 2026, 10, {"16-10": [("Santander", "Sublime", 5958.19, 5958.19)]})

    r = client.get("/saldos/mapa-saldos", params={"dia": "2026-10-16"})

    assert r.status_code == 200
    assert r.json()["contas"][0]["banco"] == "Santander"
    assert r.json()["saldo_contabilistico_total"] == 5958.19
