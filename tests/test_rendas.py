"""Mapa de Rendas cruzado com os extratos (app/services/rendas.py)."""
import os
from datetime import date, timedelta

import openpyxl

from app.db.models import MovimentoBancario
from app.services import rendas
from app.services.fluxos_conhecidos import ler_contratos_rendas
from app.services.rendas import detalhe_rendas, distribuir_pagamentos


def _pag(dia, n_meses=1):
    return {"dia": dia, "valor": 100.0 * n_meses, "n_meses": n_meses}


def test_renda_paga_no_mes_seguinte_tapa_o_mes_em_atraso():
    a, b = _pag(date(2026, 9, 2)), _pag(date(2026, 9, 28))
    pagos = distribuir_pagamentos([b, a], list(range(1, 10)))
    assert pagos[9] is a and pagos[1] is b  # o 2.º pagamento de setembro vai para o mais antigo livre


def test_varias_rendas_pagam_os_meses_em_atraso():
    pagos = distribuir_pagamentos(
        [_pag(date(2026, m, 5)) for m in range(1, 7)] + [_pag(date(2026, 9, 5), n_meses=3)], list(range(1, 10)),
    )
    assert sorted(pagos) == list(range(1, 7)) + [7, 8, 9]


def test_meses_sem_contrato_nao_recebem_pagamentos():
    pagos = distribuir_pagamentos([_pag(date(2026, 3, 1), n_meses=2)], [3, 4, 5])
    assert sorted(pagos) == [3, 4]


def _mapa_rendas(caminho, linhas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RENDAS"
    ws.cell(1, 62, 2026)  # bloco de 2026 em BJ, como no Mapa real
    for i, (empresa, cliente, renda, marcas) in enumerate(linhas, start=3):
        ws.cell(i, 1, empresa)
        ws.cell(i, 3, cliente)
        ws.cell(i, 6, "Fânzeres")
        ws.cell(i, 11, "J")
        ws.cell(i, 23, renda)
        for mes, marca in marcas.items():
            ws.cell(i, 61 + mes, marca)
    wb.save(caminho)


def test_le_as_marcas_do_ano(tmp_path):
    caminho = tmp_path / "rendas.xlsx"
    _mapa_rendas(caminho, [("H.C.N.", "Maria Santos", 50, {1: "X", 2: "x", 3: "-"})])
    [c] = ler_contratos_rendas(str(caminho), 2026)
    assert c.meses_mapa[1] == "X" and c.meses_mapa[2] == "X" and c.meses_mapa[3] == "-" and c.meses_mapa[4] is None


def _extratos_completos(db_session, ate):
    """Um movimento qualquer por dia, para nenhum mês contar como incompleto."""
    dia = date(2026, 1, 1)
    while dia <= ate:
        db_session.add(MovimentoBancario(dia=dia, empresa="X", descricao="OUTRO", valor=1, ficheiro_origem="x"))
        dia += timedelta(days=1)


def test_detalhe_compara_extrato_com_o_mapa(db_session, tmp_path, monkeypatch):
    caminho = tmp_path / "rendas.xlsx"
    _mapa_rendas(caminho, [("H.C.N.", "Maria de Lurdes Santos", 50, {1: "X", 2: "X"})])
    monkeypatch.setattr(rendas, "_onedrive_raiz", lambda: str(tmp_path))
    monkeypatch.setattr(rendas, "mapa_rendas_mais_recente", lambda raiz, ate: str(caminho))
    for dia in (date(2026, 1, 5), date(2026, 3, 6)):
        db_session.add(MovimentoBancario(dia=dia, empresa="HABISERVE CONSTRUCOES NORTE,LDA",
                                         descricao="TRF MARIA LURDES SANTOS", valor=50, ficheiro_origem="x"))
    _extratos_completos(db_session, date(2026, 5, 10))
    db_session.commit()

    r = detalhe_rendas(db_session, ate=date(2026, 4, 10))
    [c] = r["contratos"]
    estados = [m["estado"] for m in c["meses"][:5]]
    assert estados == ["pago", "registado no Mapa", "por registar", "por receber", None]
    assert r["resumo"]["meses_por_registar"] == 1 and r["resumo"]["meses_em_falta"] == 0

    r = detalhe_rendas(db_session, ate=date(2026, 5, 10))
    assert r["contratos"][0]["meses_em_falta"] == [4]


def test_mes_com_extratos_incompletos_nao_fica_em_falta(db_session, tmp_path, monkeypatch):
    caminho = tmp_path / "rendas.xlsx"
    _mapa_rendas(caminho, [("H.C.N.", "Maria de Lurdes Santos", 50, {})])
    monkeypatch.setattr(rendas, "_onedrive_raiz", lambda: str(tmp_path))
    monkeypatch.setattr(rendas, "mapa_rendas_mais_recente", lambda raiz, ate: str(caminho))
    _extratos_completos(db_session, date(2026, 2, 28))  # março sem extratos
    db_session.add(MovimentoBancario(dia=date(2026, 1, 5), empresa="HABISERVE CONSTRUCOES NORTE,LDA",
                                     descricao="TRF MARIA LURDES SANTOS", valor=50, ficheiro_origem="x"))
    db_session.add(MovimentoBancario(dia=date(2026, 1, 6), empresa="HABISERVE CONSTRUCOES NORTE,LDA",
                                     descricao="TRF MARIA LURDES SANTOS", valor=5000, ficheiro_origem="x"))
    db_session.commit()

    r = detalhe_rendas(db_session, ate=date(2026, 4, 10))
    [c] = r["contratos"]
    assert [m["estado"] for m in c["meses"][:4]] == ["por registar", "em falta", "sem extrato", "por receber"]
    assert r["meses_sem_extrato"] == [3] and len(c["pagamentos"]) == 1  # 100 rendas de uma vez não é renda
