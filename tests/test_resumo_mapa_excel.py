from datetime import date

import openpyxl
from openpyxl.worksheet.table import Table

from app.services import resumo_mapa_excel, resumo_saldos


def _mapa(caminho):
    """Como o Mapa de outubro: Resumo com fórmulas para o total de cada folha
    (dia 01 com =SUM, dia 02 com uma Tabela do Excel), saldos e CPCVs."""
    wb = openpyxl.Workbook()
    res = wb.active
    res.title = "Resumo"
    res["B2"], res["K2"], res["O2"], res["P2"] = "Recebimentos", "Pagamentos", "Saldo Bancário Disponível", \
        "Saldo Bancário Indisponivel"
    for col, texto in zip("BCDEFGHIJKLM", ["Previsão", "Recebido", "Desvio", "# de CPCVs", "Valor Tabela",
                                           "Valor Venda", "Valor Recebido", "# de Escrituras", "Valor Recebido",
                                           "Previsão", "Pago", "Desvio"]):
        res[f"{col}3"] = texto
    d1 = wb.create_sheet("01")
    d1["D5"], d1["D6"], d1["J5"] = 1000, 500.5, 200
    d1["D20"], d1["J20"] = "=SUM(D5:D19)", "=SUM(J5:J19)"
    d2 = wb.create_sheet("02")
    for c, texto in zip("GHIJKL", ["ref", "Descrição", "Valor Previsto", "Valor Pago", "Empresa", "Imputação"]):
        d2[f"{c}4"] = texto
    d2["J5"], d2["J6"] = 300, 50
    d2["J8"] = "=SUBTOTAL(109,Tabela1[Valor Pago])"
    tabela = Table(displayName="Tabela1", ref="G4:L8")
    tabela.totalsRowCount = 1
    d2.add_table(tabela)
    d2["D10"] = "=SUM(D5:D9)"
    res.append([])  # linha 4
    res["A5"], res["C5"], res["L5"], res["O5"], res["P5"] = 1, "=+'01'!D20", "='01'!J20", 158624.38, 42865.76
    res["E5"], res["H5"], res["I5"], res["J5"] = "VivaMonte - Fração 0.2\nAbadias - Fração B", 162500, \
        "00.ST.129.032", 38000
    res["A6"], res["C6"], res["L6"], res["O6"], res["P6"] = 2, "='02'!D10", "='02'!J8", 161824.48, 46708.30
    res["A7"], res["C7"], res["L7"] = 3, "='03'!D9", "='03'!J9"  # sábado: sem folha nem saldo
    res["A36"], res["C36"] = "Total", "=SUM(C5:C35)"
    wb.save(caminho)


def test_le_o_resumo_como_o_excel(tmp_path):
    caminho = tmp_path / "10 - Mapa.xlsx"
    _mapa(caminho)
    dias = resumo_mapa_excel.ler_resumo(str(caminho))
    assert dias[1] == {"recebido": 1500.5, "pago": 200.0, "saldo_disponivel": 158624.38, "saldo_indisponivel": 42865.76,
                       "saldo_bancario": 158624.38 + 42865.76, "cpcv": 162500.0, "n_cpcv": 2, "escritura": 38000.0,
                       "n_escritura": 1}
    assert (dias[2]["recebido"], dias[2]["pago"]) == (0.0, 350.0)  # coluna da Tabela
    assert (dias[3]["recebido"], dias[3]["saldo_disponivel"]) == (None, None)  # folha 03 não existe
    assert 36 not in dias


def test_resumo_mensal_usa_o_mapa_do_mes(client, db_session, monkeypatch):
    """Num mês com Mapa, recebido/pago/saldos/CPCVs são os do Resumo do Mapa
    (os valores das células), não os dos extratos."""
    from app.db.models import MovimentoBancario, SaldoDiario

    db_session.add_all([
        SaldoDiario(dia=date(2026, 10, 1), entidade="A,LDA", saldo_contabilistico=1, saldo_disponivel=1),
        SaldoDiario(dia=date(2026, 10, 2), entidade="A,LDA", saldo_contabilistico=1, saldo_disponivel=1),
        MovimentoBancario(dia=date(2026, 10, 1), empresa="A,LDA", descricao="x", valor=999999, ficheiro_origem="x"),
    ])
    db_session.commit()
    mapa = {
        date(2026, 10, 1): {"recebido": 1500.5, "pago": 200.0, "saldo_disponivel": 100.0, "saldo_indisponivel": 20.0,
                            "saldo_bancario": 120.0, "cpcv": 162500.0, "n_cpcv": 2, "escritura": 38000.0,
                            "n_escritura": 1},
        date(2026, 10, 2): {"recebido": 0.0, "pago": 350.0, "saldo_disponivel": None, "saldo_indisponivel": None,
                            "saldo_bancario": None, "cpcv": 0.0, "n_cpcv": 0, "escritura": 0.0, "n_escritura": 0},
    }
    monkeypatch.setattr(resumo_saldos, "resumo_do_mes", lambda ano, mes: mapa if (ano, mes) == (2026, 10) else None)
    r = client.get("/saldos/resumo-mensal", params={"ano": 2026}).json()
    outubro = [m for m in r["meses"] if m["mes"] == 10][0]
    assert outubro["fonte"] == "mapa"
    assert (outubro["recebido"], outubro["pago"], outubro["cpcv"], outubro["n_cpcv"], outubro["escritura"]) == (
        1500.5, 550.0, 162500.0, 2, 38000.0)
    assert (outubro["saldo_fim"], outubro["saldo_contabilistico_fim"]) == (100.0, 120.0)
    diario = {d["dia"]: d for d in r["diario"]}
    assert (diario["2026-10-01"]["recebido"], diario["2026-10-01"]["liquidez"]) == (1500.5, 100.0)
    # dia sem saldo na linha do Resumo: o do dia antes
    assert (diario["2026-10-02"]["pago"], diario["2026-10-02"]["liquidez"]) == (350.0, 100.0)


def test_le_o_resumo_antigo_de_janeiro_a_maio(tmp_path):
    """Jan-mai: cabeçalho nas linhas 3/4, Total RECEBIDO (= Recebido +
    Excepcional), Valor PAGO e um só Saldo Bancário - que às vezes dá #REF!."""
    wb = openpyxl.Workbook()
    res = wb.active
    res.title = "Resumo"
    res["B3"], res["G3"], res["L3"], res["M3"] = "Recebimentos", "Pagamentos", "LIQUIDEZ", "Saldo"
    for col, texto in zip("BCDEFGHIJKM", ["Previsão", "Recebido", "Desvio", "Excepcional", "Total RECEBIDO",
                                          "Previsão", "Pago", "Desvio", "Valor PAGO", "Compromisso", "Bancário"]):
        res[f"{col}4"] = texto
    dia = wb.create_sheet("02")
    dia["I25"], dia["D25"] = 700, 250
    res["A9"], res["C9"], res["E9"], res["F9"] = 2, "='02'!I25", 50, "=C9+E9"
    res["H9"], res["J9"], res["M9"] = "='02'!D25", "=H9", "=#REF!+F14-J14"
    res["A10"], res["M10"] = 3, 1234.5
    caminho = tmp_path / "03 - Mapa.xlsx"
    wb.save(caminho)
    dias = resumo_mapa_excel.ler_resumo(str(caminho))
    assert (dias[2]["recebido"], dias[2]["pago"]) == (750.0, 250.0)
    assert (dias[2]["saldo_disponivel"], dias[2]["saldo_bancario"]) == (None, None)  # #REF!: sem valor
    assert dias[3]["saldo_bancario"] == 1234.5
    assert (dias[2]["cpcv"], dias[2]["n_cpcv"]) == (None, None)  # esta forma do Resumo não tem CPCVs


def test_valores_escritos_como_formula_nas_linhas_contam(tmp_path):
    """Nas folhas antigas há valores escritos à mão como fórmula ("=5.92+3.95")
    e totais com SUBTOTAL - o total tem de ser o do Excel."""
    wb = openpyxl.Workbook()
    res = wb.active
    res.title = "Resumo"
    res["B2"], res["K2"] = "Recebimentos", "Pagamentos"
    res["B3"], res["C3"], res["K3"], res["L3"] = "Previsão", "Recebido", "Previsão", "Pago"
    dia = wb.create_sheet("05")
    dia["D5"], dia["D6"], dia["D7"] = 100, "=5.92+3.95", "=D5*2"
    dia["D20"] = "=SUBTOTAL(9,D5:D19)"
    dia["J5"], dia["J6"] = "=10+0.5", "#texto"
    dia["J20"] = "=SUM(SUM(J5:J6),1)"
    res["A5"], res["C5"], res["L5"] = 5, "='05'!D20", "='05'!J20"
    res["A6"], res["C6"] = 6, "='06'!D20+#REF!"
    caminho = tmp_path / "mapa.xlsx"
    wb.save(caminho)
    dias = resumo_mapa_excel.ler_resumo(str(caminho))
    assert round(dias[5]["recebido"], 2) == 309.87  # 100 + 9.87 + 200
    assert dias[5]["pago"] == 11.5
    assert dias[6]["recebido"] is None  # erro no Mapa: sem valor, nunca um total errado


def test_movimentos_do_extrato_mensal_e_dos_diarios_contam_uma_vez(client, db_session, monkeypatch):
    """Jan-mar só tinham extratos mensais; a 20/10/2026 passaram a ter também
    pastas diárias e os mesmos movimentos entraram duas vezes na BD. Não se
    apagam: os dos ficheiros diários de um dia/empresa com extrato mensal
    não contam."""
    from app.db.models import MovimentoBancario, SaldoDiario

    monkeypatch.setattr(resumo_saldos, "resumo_do_mes", lambda ano, mes: None)
    db_session.add_all([
        SaldoDiario(dia=date(2026, 2, 3), entidade="A,LDA", saldo_contabilistico=1, saldo_disponivel=1),
        MovimentoBancario(dia=date(2026, 2, 3), empresa="A,LDA", descricao="TRF", valor=100,
                          ficheiro_origem="mensal:A,LDA.xlsx"),
        MovimentoBancario(dia=date(2026, 2, 3), empresa="A,LDA", descricao="EDP", valor=-40,
                          ficheiro_origem="mensal:A,LDA.xlsx"),
        MovimentoBancario(dia=date(2026, 2, 3), empresa="A,LDA", descricao="TRF", valor=100,
                          ficheiro_origem="03-02-2026_A,LDA.xlsx"),  # o mesmo, outra vez
        MovimentoBancario(dia=date(2026, 2, 3), empresa="B,LDA", descricao="TRF", valor=7,
                          ficheiro_origem="03-02-2026_B,LDA.xlsx"),  # B sem extrato mensal: conta
    ])
    db_session.commit()
    r = client.get("/saldos/resumo-mensal", params={"ano": 2026}).json()
    fevereiro = [m for m in r["meses"] if m["mes"] == 2][0]
    assert (fevereiro["recebido"], fevereiro["pago"]) == (107.0, 40.0)
    assert {d["dia"]: d for d in r["diario"]}["2026-02-03"]["recebido"] == 107.0
    assert db_session.query(MovimentoBancario).count() == 4  # nada apagado


def test_importador_nao_volta_a_importar_um_dia_que_veio_do_extrato_mensal(tmp_path, db_session):
    from app.db.models import MovimentoBancario
    from app.services import onedrive_sync

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Data mov. ", "Data-valor ", "Descrição ", "Montante ", "Saldo contabilístico após movimento "])
    ws.append(["03-02-2026", "03-02-2026", "TRF ", "100,00", "1.100,00"])
    for empresa in ("A,LDA", "B,LDA"):
        wb.save(tmp_path / f"03-02-2026_{empresa}.xlsx")
    db_session.add(MovimentoBancario(dia=date(2026, 2, 3), empresa="A,LDA", descricao="TRF", valor=100,
                                     ficheiro_origem="mensal:A,LDA.xlsx"))
    db_session.commit()
    assert onedrive_sync._importar_movimentos_em_falta(db_session, str(tmp_path / "03-02-2026_A,LDA.xlsx"),
                                                       date(2026, 2, 3)) == 0
    assert onedrive_sync._importar_movimentos_em_falta(db_session, str(tmp_path / "03-02-2026_B,LDA.xlsx"),
                                                       date(2026, 2, 3)) == 1
    assert db_session.query(MovimentoBancario).count() == 2


def test_avisa_celulas_do_resumo_que_nao_leem_o_total_do_dia(tmp_path):
    """21/10/2026, Resumo de setembro: o dia 8 lia '08'!D31 com o total já
    em D37 (contava 0); em julho o dia 4 lia a folha do dia 2."""
    wb = openpyxl.Workbook()
    res = wb.active
    res.title = "Resumo"
    res["B2"], res["K2"] = "Recebimentos", "Pagamentos"  # formato atual: Recebido em C, Pago em L
    res["B3"], res["C3"], res["K3"], res["L3"] = "Previsão", "Recebido", "Previsão", "Pago"
    for nome, total in (("04", 30), ("08", 37), ("09", 20)):
        folha = wb.create_sheet(nome)
        folha[f"D{total}"], folha[f"J{total}"] = f"=SUM(D5:D{total - 1})", f"=SUM(J5:J{total - 1})"
    wb.create_sheet("02")["D52"] = "=SUM(D5:D51)"
    res["A9"], res["C9"] = 4, "=+'02'!D52"                      # outra folha
    res["A13"], res["C13"], res["L13"] = 8, "='08'!D31", "='08'!J37"  # C: linha errada; L: certa
    res["A14"], res["C14"], res["L14"] = 9, "='09'!D20", "='09'!J20"  # certas
    caminho = tmp_path / "09 - Mapa.xlsx"
    wb.save(caminho)

    avisos = resumo_mapa_excel.referencias_desalinhadas(str(caminho))

    assert avisos == ["dia 4: C9 lê a folha '02' (=+'02'!D52) em vez da folha '04'",
                      "dia 8: C13 lê D31 da folha '08', mas o total do dia está em D37"]


def test_resumo_no_formato_antigo_nao_da_avisos(tmp_path):
    """Jan-jun: folhas dos dias com outra disposição e vários subtotais (o
    total do dia 9 de janeiro é J32, com um =SUM(I30) solto em J30) - os
    avisos eram falsos alarmes."""
    wb = openpyxl.Workbook()
    res = wb.active
    res.title = "Resumo"
    res["B3"], res["G3"] = "Recebimentos", "Pagamentos"
    for col, texto in zip("BCDEFGHIJ", ["Previsão", "Recebido", "Desvio", "Excepcional", "Total RECEBIDO",
                                        "Previsão", "Pago", "Desvio", "Valor PAGO"]):
        res[f"{col}4"] = texto
    dia = wb.create_sheet("09")
    dia["J30"], dia["J32"] = "=SUM(I30)", "=SUM(J5:J30)"
    res["A14"], res["C14"] = 9, "='09'!J32"
    caminho = tmp_path / "01 - Mapa.xlsx"
    wb.save(caminho)
    assert resumo_mapa_excel.referencias_desalinhadas(str(caminho)) == []
