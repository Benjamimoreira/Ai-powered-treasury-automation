from datetime import date

import openpyxl

from app.db.models import LinhaMapa, MovimentoBancario, Reconciliacao
from app.services import atualizacao_dados
from app.services.mapa_importer import importar_dia_do_mapa, sincronizar_dia_do_mapa

DIA = date(2026, 10, 9)


def _mapa(tmp_path, recebimentos, pagamentos=()):
    """Folha "09" com recebimentos (B-F) e pagamentos (H-L) a partir da
    linha 5 e a linha de totais (=SUM) logo a seguir."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "09"
    for i, (desc, previsto, real, empresa, imputacao) in enumerate(recebimentos):
        r = 5 + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"F{r}"] = desc, previsto, real, empresa, imputacao
    for i, (desc, previsto, real, empresa, imputacao) in enumerate(pagamentos):
        r = 5 + i
        ws[f"H{r}"], ws[f"I{r}"], ws[f"J{r}"], ws[f"K{r}"], ws[f"L{r}"] = desc, previsto, real, empresa, imputacao
    ws["C20"] = "=SUM(C5:C19)"
    caminho = tmp_path / "mapa.xlsx"
    wb.save(caminho)
    return str(caminho)


def test_sincronizar_dia_do_mapa_acompanha_a_folha_sem_mexer_no_reconciliado(tmp_path, db_session):
    primeira = _mapa(tmp_path, [
        ("CPCV - 00.PO.23.035", 41750, None, "H.I.I", "CPCV"),
        ("Renda", 500, 500, "HCN", "Rendas"),
        ("Sai da folha", 10, None, "HCN", None),
    ])
    importar_dia_do_mapa(db_session, primeira, DIA)
    db_session.commit()
    renda = db_session.query(LinhaMapa).filter_by(linha=6).one()
    movimento = MovimentoBancario(dia=DIA, empresa="HCN", descricao="TRF RENDA", valor=500, ficheiro_origem="x.xlsx")
    db_session.add(movimento)
    db_session.flush()
    db_session.add(Reconciliacao(movimento_id=movimento.id, linha_id=renda.id, tipo_match="exato"))
    db_session.commit()

    # mais tarde no mesmo dia: o CPCV foi recebido, a renda mudou de valor
    # (mas está reconciliada), uma linha saiu e entrou um pagamento novo
    (tmp_path / "mapa.xlsx").unlink()
    segunda = _mapa(tmp_path, [
        ("CPCV - 00.PO.23.035", 41750, 41750, "H.I.I", "CPCV"),
        ("Renda", 999, 999, "HCN", "Rendas"),
    ], pagamentos=[("Fornecedor", 200, None, "HCN", "Obras")])

    assert sincronizar_dia_do_mapa(db_session, segunda, DIA) == 3
    db_session.commit()
    linhas = {(l.tipo, l.linha): l for l in db_session.query(LinhaMapa).filter_by(dia=DIA).all()}
    assert linhas[("recebimento", 5)].pago == 41750
    assert linhas[("recebimento", 6)].pago == 500  # reconciliada: não mexe
    assert ("recebimento", 7) not in linhas
    assert linhas[("pagamento", 5)].previsto == -200
    # repetir sem mudanças não faz nada
    assert sincronizar_dia_do_mapa(db_session, segunda, DIA) == 0


def test_fim_de_script_que_produz_dados_importa_hoje_e_ontem(monkeypatch, session_factory):
    importados = []

    def importar_falso(db, dia):
        importados.append(dia)
        return {"dias_com_movimentos_novos": [dia.isoformat()], "dias_com_saldos_novos": [],
                "dias_com_mapa_novo": [], "erros": []}

    monkeypatch.setattr("app.services.onedrive_sync.atualizar_dados_do_dia", importar_falso)
    monkeypatch.setattr("app.db.session.SessionLocal", session_factory)

    atualizacao_dados.sincronizar_apos_script("extrair_faturas", "ok")
    atualizacao_dados.sincronizar_apos_script("preencher_mapa", "erro")
    assert importados == []

    antes = atualizacao_dados.estado()["atualizado_em"]
    atualizacao_dados.sincronizar_apos_script("movimentos_dia_atual_cgd", "warning")
    assert len(importados) == 2 and importados[1] - importados[0] == date(2026, 1, 2) - date(2026, 1, 1)
    estado = atualizacao_dados.estado()
    assert estado["origem"] == "fim de movimentos_dia_atual_cgd"
    assert estado["dias"] == sorted(d.isoformat() for d in importados)
    assert estado["atualizado_em"] >= antes


def test_executar_dispara_a_sincronizacao_e_sync_estado_responde(client, monkeypatch):
    chamadas = []
    monkeypatch.setattr("app.routers.monitorizacao.sincronizar_apos_script",
                        lambda script, status: chamadas.append((script, status)))
    resposta = client.post("/monitorizacao/scripts/preencher_mapa/executar", json={"status": "ok"})
    assert resposta.status_code == 200
    assert chamadas == [("preencher_mapa", "ok")]

    atualizacao_dados.registar_resultado({"dias_com_mapa_novo": ["2026-10-09"]}, "teste")
    corpo = client.get("/sync/estado").json()
    assert corpo["origem"] == "teste" and corpo["dias"] == ["2026-10-09"]
    # sem nada de novo, o estado não muda
    assert atualizacao_dados.registar_resultado({"dias_com_mapa_novo": []}, "vazio") is False
    assert client.get("/sync/estado").json()["origem"] == "teste"


def test_inicio_fim_e_agendamento_no_estado_dos_scripts(client):
    client.post("/monitorizacao/scripts/movimentos_diarios_cgd/inicio", json={"origem": "teste"})
    client.post("/monitorizacao/tarefas-agendadas", json={"tarefas": [
        {"nome": "Movimentos Diários CGD", "estado": "Ready", "proxima_execucao": "2026-10-13T07:30:00.0000000Z"},
        {"nome": "Tesouraria - Enviar Mapa Diario", "estado": "Disabled", "proxima_execucao": "2026-10-12T15:30:00Z"},
    ]})
    scripts = {s["nome"]: s for s in client.get("/monitorizacao/scripts").json()["scripts"]}
    diarios = scripts["movimentos_diarios_cgd"]
    assert diarios["a_correr_desde"] is not None
    assert diarios["agendamento"] == "ativo" and diarios["proxima_execucao"] == "2026-10-13T07:30:00Z"
    assert scripts["enviar_mapa_smtp"]["agendamento"] == "desativado"
    assert scripts["deploy_scripts_cgd"]["agendamento"] == "a pedido"

    # o fim tira-o de "a correr"
    client.post("/monitorizacao/scripts/movimentos_diarios_cgd/executar", json={"status": "ok"})
    scripts = {s["nome"]: s for s in client.get("/monitorizacao/scripts").json()["scripts"]}
    assert scripts["movimentos_diarios_cgd"]["a_correr_desde"] is None
    assert [t["nome"] for t in client.get("/monitorizacao/tarefas-agendadas").json()["tarefas"]] == [
        "Movimentos Diários CGD", "Tesouraria - Enviar Mapa Diario"]


def test_fim_da_extracao_pede_o_preencher_mapa_exceto_se_ja_estiver_a_correr(monkeypatch, session_factory):
    monkeypatch.setattr("app.services.onedrive_sync.atualizar_dados_do_dia",
                        lambda db, dia: {"dias_com_movimentos_novos": [], "dias_com_saldos_novos": [],
                                         "dias_com_mapa_novo": [], "erros": []})
    monkeypatch.setattr("app.db.session.SessionLocal", session_factory)
    monkeypatch.delenv("SCRIPTS_PREENCHIMENTO_RAIZ", raising=False)
    monkeypatch.delenv("SCRIPTS_LOG_DIR", raising=False)
    from app.db.models import PedidoCorrida
    from app.services.monitorizacao import registar_inicio

    atualizacao_dados.sincronizar_apos_script("movimentos_dia_atual_cgd", "warning")
    db = session_factory()
    assert [p.script for p in db.query(PedidoCorrida).all()] == ["preencher_mapa"]
    db.query(PedidoCorrida).delete(); db.commit()

    registar_inicio(db, "preencher_mapa", origem="teste")  # já a correr: não pede outro
    atualizacao_dados.sincronizar_apos_script("movimentos_diarios_cgd", "ok")
    assert db.query(PedidoCorrida).count() == 0
    atualizacao_dados.sincronizar_apos_script("extrair_faturas", "ok")  # não é extração dos extratos
    assert db.query(PedidoCorrida).count() == 0
    db.close()
