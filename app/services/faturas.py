"""Faturas recebidas em faturas@vidor.pt, reportadas pelo
recolher_faturas_recebidas.py (pasta "Fornecedores", fora do Docker) via
POST /faturas/recebidas - guarda na BD o mesmo que já vai para o Excel
mensal, para a dashboard mostrar sem precisar de abrir o ficheiro.

Desde 09/10/2026 o script deixou de reportar à API (continua a escrever o
Excel e os PDFs no SharePoint), por isso a API também as vai buscar ao
Excel mensal: importar_faturas_excel (botão na Monitorização e o
sincronizador de 15 em 15 min)."""
import os
import time
from datetime import date, datetime
from typing import List, Optional
from urllib.parse import unquote

from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.db.models import FaturaRecebida
from app.models import FaturaRecebidaIn

# Pasta do SharePoint (dentro de ONEDRIVE_RAIZ) onde o recolher_faturas_recebidas.py
# grava, por ano, o Excel mensal "MM-Faturas Recebidas.xlsx" e os PDFs em
# "Faturas Recebidas\AAAA-MM\". O pdf_relativo é relativo à pasta do ano.
PASTA_FATURAS = os.path.join("DOCUMENTOS A TRATAR", "AFaturas")
NOME_SCRIPT_IMPORTACAO = "importar_faturas_recebidas"


def pasta_faturas_do_ano(ano: int) -> Optional[str]:
    raiz = os.environ.get("ONEDRIVE_RAIZ")
    return os.path.join(raiz, PASTA_FATURAS, str(ano)) if raiz else None


def caminho_excel_faturas(ano: int, mes: int) -> Optional[str]:
    pasta = pasta_faturas_do_ano(ano)
    return os.path.join(pasta, f"{mes:02d}-Faturas Recebidas.xlsx") if pasta else None


def _texto(valor) -> Optional[str]:
    if valor is None or valor == "":
        return None
    return str(valor)


def ler_excel_faturas(caminho: str, ano: int) -> List[FaturaRecebidaIn]:
    """Linhas da folha "Todas" do Excel mensal (colunas dia, hora, remetente,
    ..., valor_fatura, hiperligacao_pdf, id). O dia vem como "DD-MM" (o ano é
    o da pasta) e o PDF é a hiperligação da célula "Abrir PDF", relativa à
    pasta do ano e com %20 & cia."""
    import openpyxl

    # sem read_only: as hiperligações só são lidas no modo normal
    livro = openpyxl.load_workbook(caminho, data_only=True)
    folha = livro["Todas"] if "Todas" in livro.sheetnames else livro.worksheets[0]
    cabecalho = [str(c.value).strip().lower() if c.value is not None else "" for c in folha[1]]
    coluna = {nome: i for i, nome in enumerate(cabecalho) if nome}

    linhas = []
    for celulas in folha.iter_rows(min_row=2):
        def valor(nome):
            i = coluna.get(nome)
            return celulas[i].value if i is not None else None

        outlook_id = _texto(valor("id"))
        dia = valor("dia")
        if not outlook_id or not dia:
            continue
        if isinstance(dia, datetime):
            dia = dia.date()
        if isinstance(dia, date):
            data = datetime.combine(dia, datetime.min.time())
        else:
            dd, mm = str(dia).strip().split("-")[:2]
            data = datetime(ano, int(mm), int(dd))
        hora = valor("hora")
        if hora:
            hh, mi = str(hora).strip().split(":")[:2]
            data = data.replace(hour=int(hh), minute=int(mi))

        pdf_relativo = None
        i_pdf = coluna.get("hiperligacao_pdf")
        if i_pdf is not None and celulas[i_pdf].hyperlink and celulas[i_pdf].hyperlink.target:
            pdf_relativo = unquote(celulas[i_pdf].hyperlink.target).replace("/", "\\")

        n_anexos = valor("n_anexos_pdf")
        linhas.append(FaturaRecebidaIn(
            outlook_id=outlook_id,
            data_recebido=data,
            remetente=_texto(valor("remetente")),
            assunto=_texto(valor("assunto")),
            motivo=_texto(valor("motivo")),
            empresa=_texto(valor("empresa")),
            fornecedor=_texto(valor("fornecedor")),
            nif_fornecedor=_texto(valor("nif_fornecedor")),
            n_anexos_pdf=int(n_anexos) if isinstance(n_anexos, (int, float)) else None,
            debito=_texto(valor("debito")),
            credito=_texto(valor("credito")),
            saldo=_texto(valor("saldo")),
            valor_fatura=_texto(valor("valor_fatura")),
            pdf_relativo=pdf_relativo,
        ))
    return linhas


def meses_a_importar(hoje: date) -> List[tuple]:
    """Mês corrente e, nos primeiros 5 dias, também o anterior (as faturas
    do fim do mês podem chegar ao Excel depois da meia-noite)."""
    meses = [(hoje.year, hoje.month)]
    if hoje.day <= 5:
        meses.insert(0, (hoje.year - 1, 12) if hoje.month == 1 else (hoje.year, hoje.month - 1))
    return meses


def importar_faturas_excel(db: Session, meses: Optional[List[tuple]] = None) -> dict:
    """Importa para a BD as faturas do(s) Excel mensal(ais) que ainda lá não
    estão (pelo id do Outlook). Devolve {"novas", "duplicadas", "linhas", "erros"}."""
    meses = meses or meses_a_importar(date.today())
    novas = duplicadas = 0
    linhas_log, erros = [], []
    for ano, mes in meses:
        caminho = caminho_excel_faturas(ano, mes)
        if not caminho:
            erros.append("ONEDRIVE_RAIZ não configurado.")
            break
        if not os.path.isfile(caminho):
            erros.append(f"Excel não encontrado: {os.path.basename(caminho)} ({ano})")
            continue
        try:
            linhas = ler_excel_faturas(caminho, ano)
        except Exception as exc:  # ficheiro a meio de sincronizar, aberto no Excel, ...
            erros.append(f"{os.path.basename(caminho)}: {exc}")
            continue
        resultado = registar_faturas(db, linhas)
        novas += resultado["novas"]
        duplicadas += resultado["duplicadas"]
        linhas_log.append(f"{mes:02d}/{ano}: {len(linhas)} no Excel, {resultado['novas']} novas")
    return {"novas": novas, "duplicadas": duplicadas, "linhas": linhas_log, "erros": erros}


def importar_e_reportar(db: Session, meses: Optional[List[tuple]] = None, reportar_sem_novas: bool = True) -> dict:
    """importar_faturas_excel + registo da corrida na Monitorização
    (importar_faturas_recebidas). Sem novas nem erros só regista se
    reportar_sem_novas - o sincronizador não, para não encher o histórico."""
    from app.services.monitorizacao import registar_execucao

    inicio = time.monotonic()
    resultado = importar_faturas_excel(db, meses)
    if resultado["erros"] or resultado["novas"] or reportar_sem_novas:
        registar_execucao(
            db, NOME_SCRIPT_IMPORTACAO,
            "warning" if resultado["erros"] else "ok",
            erro="; ".join(resultado["erros"]) or None,
            log=resultado["linhas"] + resultado["erros"],
            duracao_segundos=time.monotonic() - inicio,
        )
    return resultado


def caminho_pdf(fatura: FaturaRecebida, fornecedores_raiz: Optional[str] = None) -> Optional[str]:
    """Caminho do PDF da fatura, ou None. Procura na pasta do ano no
    SharePoint (onde o script grava hoje) e, para faturas antigas, em
    FORNECEDORES_RAIZ (onde gravava antes). Nunca sai da pasta de cada raiz."""
    if not fatura.pdf_relativo:
        return None
    # pdf_relativo é gravado com barras invertidas (Windows, de onde vem o
    # script) - num container Linux "\" não é separador de caminho
    relativo = fatura.pdf_relativo.replace("\\", "/")
    for raiz in filter(None, [pasta_faturas_do_ano(fatura.dia.year), fornecedores_raiz]):
        raiz = os.path.realpath(raiz)
        caminho = os.path.realpath(os.path.join(raiz, relativo))
        # defesa contra pdf_relativo malformado ("..\..")
        if os.path.commonpath([raiz, caminho]) == raiz and os.path.isfile(caminho):
            return caminho
    return None


def registar_faturas(db: Session, linhas: List[FaturaRecebidaIn]) -> dict:
    novas = duplicadas = 0
    vistos_neste_pedido = set()
    for linha in linhas:
        # duas verificações: já na BD (pedidos anteriores) e já visto neste
        # mesmo pedido (o mesmo email pode aparecer 2x no lote, ex. num
        # backfill a partir do Excel - sem isto, a 2ª linha só falhava na
        # BD com "unique constraint" em vez de ser tratada como duplicada)
        if linha.outlook_id in vistos_neste_pedido:
            duplicadas += 1
            continue
        existe = db.query(FaturaRecebida).filter(FaturaRecebida.outlook_id == linha.outlook_id).first()
        if existe:
            duplicadas += 1
            continue
        vistos_neste_pedido.add(linha.outlook_id)
        db.add(FaturaRecebida(
            outlook_id=linha.outlook_id,
            dia=linha.data_recebido.date(),
            hora=linha.data_recebido.strftime("%H:%M"),
            remetente=linha.remetente,
            assunto=linha.assunto,
            motivo=linha.motivo,
            empresa=linha.empresa,
            fornecedor=linha.fornecedor,
            nif_fornecedor=linha.nif_fornecedor,
            n_anexos_pdf=linha.n_anexos_pdf,
            debito=linha.debito,
            credito=linha.credito,
            saldo=linha.saldo,
            valor_fatura=linha.valor_fatura,
            pdf_relativo=linha.pdf_relativo,
        ))
        novas += 1
    db.commit()
    return {"novas": novas, "duplicadas": duplicadas}


def obter_fatura(db: Session, fatura_id: int) -> Optional[FaturaRecebida]:
    return db.query(FaturaRecebida).filter(FaturaRecebida.id == fatura_id).first()


def listar_faturas(
    db: Session,
    dia: Optional[date] = None,
    desde: Optional[date] = None,
    ate: Optional[date] = None,
    pesquisa: Optional[str] = None,
    limit: int = 200,
) -> list:
    query = db.query(FaturaRecebida)
    if dia:
        query = query.filter(FaturaRecebida.dia == dia)
    if desde:
        query = query.filter(FaturaRecebida.dia >= desde)
    if ate:
        query = query.filter(FaturaRecebida.dia <= ate)
    if pesquisa:
        # pesquisa em qualquer dia (não só na janela recente do "limit") -
        # por isso é um filtro à parte, não um contains() sobre o que já
        # foi carregado no browser
        termo = f"%{pesquisa}%"
        query = query.filter(
            or_(
                FaturaRecebida.empresa.ilike(termo),
                FaturaRecebida.fornecedor.ilike(termo),
                FaturaRecebida.nif_fornecedor.ilike(termo),
                FaturaRecebida.assunto.ilike(termo),
                FaturaRecebida.remetente.ilike(termo),
                FaturaRecebida.valor_fatura.ilike(termo),
            )
        )
    return (
        query.order_by(FaturaRecebida.dia.desc(), FaturaRecebida.hora.desc())
        .limit(limit)
        .all()
    )


# ---------------------------------------------------------------------------
# Fornecedor normalizado
# ---------------------------------------------------------------------------
# O fornecedor vem do recolher_faturas_recebidas.py (texto extraído do PDF/
# email) e em set/2026 vinha vazio em ~40% das faturas e com lixo noutras:
# "do titular IBAN", "Payment method Paid", moradas, "Referente a Factura nº
# ... | PCI ...", ou o próprio CLIENTE (uma empresa do grupo - ex. a morada
# da Palavradicional nas faturas da Miio). Aqui escolhe-se o melhor nome por
# esta ordem, sem tocar no valor original:
#   1. texto extraído, limpo, se não for lixo, morada nem empresa do grupo;
#   2. o nome mais comum para o mesmo NIF entre as faturas com texto válido;
#   3. o domínio do remetente (faturaedp@edp.pt -> "EDP");
#   4. emails internos (@vidor.pt) sem mais nada -> "(encaminhado internamente)".
import re
import unicodedata
from collections import Counter

ENCAMINHADO_INTERNAMENTE = "(encaminhado internamente)"
DOMINIOS_INTERNOS = {"vidor.pt"}
DOMINIOS_GENERICOS = {"gmail.com", "hotmail.com", "outlook.com", "outlook.pt", "sapo.pt", "live.com", "yahoo.com", "icloud.com"}
# nomes de marca com grafia própria (o resto fica com a 1.ª letra maiúscula)
NOMES_DOMINIO = {"edp": "EDP", "prio": "PRIO Energy", "galp": "Galp", "miio": "Miio", "meo": "MEO", "nos": "NOS",
                 "aquamatrix": "Águas (Aquamatrix)", "claranet": "Claranet", "ageas": "Ageas", "via-verde": "Via Verde",
                 "facebookmail": "Facebook"}
_LIXO = re.compile(
    r"^(do titular|payment method|documento|consultar|pro ?forma|referente|fatura|factura|nota de|recibo|"
    r"rua |r\. |av\.|avenida|praça|praca|largo|travessa|estrada|nif|iban|total|data|morada)",
    re.IGNORECASE,
)
_CODIGO_POSTAL = re.compile(r"\b\d{4}-\d{3}\b")
_PREFIXOS = re.compile(r"^(referente a fa[c]?tura n[ºo°]?\s*\S+\s*\|\s*|da entidade:\s*)", re.IGNORECASE)
# palavras genéricas que não identificam uma empresa do grupo
_GENERICAS = {"LDA", "SA", "UNIPESSOAL", "SOCIEDADE", "INVESTIMENTOS", "IMOBILIARIOS", "IMOBILIARIA", "CONSTRUCOES",
              "CONSTRUCAO", "GESTORA", "PARTICIPACOES", "SOCIAIS", "CIVIL", "INDUSTRIA", "NORTE", "CENTRO", "SUL",
              "ADMINISTRACAO", "CONDOMINIOS", "TURISTICOS", "SERVICOS", "EMPRESA", "CAPITAL", "RISCO", "PREDIO",
              "TRABALHO", "TEMPORARIO", "FORMACAO", "PROFISSIONAL"}


def _maiusculas(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto or "") if not unicodedata.combining(c)).upper()


def palavras_do_grupo(empresas_grupo: list) -> set:
    """A 1.ª palavra distintiva de cada empresa do grupo (HABISERVE,
    PALAVRADICIONAL, VIDOR, CERRO...) - para não confundir o cliente com o
    fornecedor. Só a primeira: as seguintes são muitas vezes apelidos
    comuns ("CORREIA" de Felizardo Correia) que apanhariam fornecedores."""
    palavras = set()
    for nome in empresas_grupo:
        distintivas = [p for p in re.findall(r"[A-Z]{5,}", _maiusculas(nome)) if p not in _GENERICAS]
        if distintivas:
            palavras.add(distintivas[0])
    return palavras


def limpar_fornecedor(texto: Optional[str], palavras_grupo: set) -> Optional[str]:
    """O texto extraído, limpo - ou None se for lixo, morada ou o cliente."""
    if not texto:
        return None
    t = _PREFIXOS.sub("", texto.strip())
    t = t.split(" | ")[0].strip(" ,;-|")
    if len(t) < 3 or _LIXO.match(t) or _CODIGO_POSTAL.search(t) or not re.search(r"[A-Za-zÀ-ÿ]{3,}", t):
        return None
    if set(re.findall(r"[A-Z]{5,}", _maiusculas(t))) & palavras_grupo:
        return None  # é uma empresa do grupo: o cliente, não o fornecedor
    return t


def fornecedor_do_remetente(remetente: Optional[str]) -> Optional[str]:
    if not remetente or "@" not in remetente:
        return None
    dominio = remetente.rsplit("@", 1)[1].strip().lower().rstrip(">")
    if dominio in DOMINIOS_INTERNOS:
        return ENCAMINHADO_INTERNAMENTE
    if dominio in DOMINIOS_GENERICOS:
        return None
    partes = dominio.split(".")
    # o nome da organização é o penúltimo bloco (hello.galp.com -> galp)
    rotulo = partes[-2] if len(partes) >= 2 else partes[0]
    return NOMES_DOMINIO.get(rotulo, rotulo.replace("-", " ").title())


def _chaves_nif(nif: Optional[str]) -> list:
    return re.findall(r"\d{9}", nif or "")


def normalizar_fornecedores(faturas: list, todas: list, empresas_grupo: list) -> list:
    """[(nome, fonte)] para cada fatura de `faturas`, com fonte "extraido" |
    "nif" | "remetente" | None. `todas` (todas as faturas da BD) serve para
    aprender o nome de cada NIF."""
    palavras_grupo = palavras_do_grupo(empresas_grupo)
    nomes_por_nif = {}
    for f in todas:
        nome = limpar_fornecedor(f.fornecedor, palavras_grupo)
        if nome:
            for nif in _chaves_nif(f.nif_fornecedor):
                nomes_por_nif.setdefault(nif, Counter())[nome] += 1

    resultado = []
    for f in faturas:
        nome = limpar_fornecedor(f.fornecedor, palavras_grupo)
        if nome:
            resultado.append((nome, "extraido"))
            continue
        do_nif = next((nomes_por_nif[n].most_common(1)[0][0] for n in _chaves_nif(f.nif_fornecedor) if n in nomes_por_nif), None)
        if do_nif:
            resultado.append((do_nif, "nif"))
            continue
        do_remetente = fornecedor_do_remetente(f.remetente)
        resultado.append((do_remetente, "remetente" if do_remetente else None))
    return resultado
