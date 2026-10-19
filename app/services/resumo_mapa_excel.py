"""Folha "Resumo" do Mapa de Pagamentos e Recebimentos de um mês, dia a dia
(pedido 19/10/2026: o resumo mensal do dashboard tem de bater com o Mapa
desse mês - recebido, pago, saldo bancário, saldo bancário disponível,
CPCVs e escrituras).

Colunas do Resumo desde julho (linha 3 é o cabeçalho, um dia por linha):
  A dia | B previsão receb. | C Recebido | E CPCVs (um por linha) |
  F Valor Tabela | G Valor Venda | H Valor Recebido (CPCVs) |
  I Escrituras (refs) | J Valor Recebido (escrituras) | K previsão pag. |
  L Pago | O Saldo Bancário Disponível | P Saldo Bancário Indisponível
Saldo bancário = disponível + indisponível. Os Mapas de janeiro a junho
têm outras formas - ver colunas_do_resumo, que encontra as colunas pelo
nome. Um valor que dá erro no Mapa (ex. #REF!) fica None.

C e L são fórmulas para o total da folha do dia (='06'!D95, que é
=SUM(D5:D94) ou =SUBTOTAL(109,Tabela1[Valor Pago])). Os scripts gravam o
Mapa com o openpyxl, que não guarda o valor calculado das fórmulas, por
isso aqui calculam-se como o Excel: só as formas que o Mapa usa (número,
referência a outra célula, SUM/SUBTOTAL de um intervalo ou coluna de
Tabela, somas de referências). Uma célula que não se sabe calcular fica
None e quem chama usa outra fonte para esse valor.
"""
import logging
import os
import re
from datetime import date
from typing import Any, Dict, Optional

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from app.services.onedrive_sync import caminho_mapa

logger = logging.getLogger(__name__)

_REF_FOLHA = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_]+))!\$?([A-Z]{1,3})\$?(\d+)")
_REF_LOCAL = re.compile(r"(?<![A-Za-z0-9_!:'])\$?([A-Z]{1,3})\$?(\d+)(?![\d(:\[])")
_FUNCAO = re.compile(r"(SUM|SUBTOTAL)\(([^()]*)\)", re.I)
_ARITMETICA = re.compile(r"^[0-9.+\-*/() eE]+$")
_CACHE: Dict[tuple, Optional[Dict[int, Dict[str, Any]]]] = {}


class _Incalculavel(Exception):
    pass


class _Calculadora:
    """Calcula células como o Excel, só para as formas que o Mapa usa: números,
    referências (noutra folha ou na mesma), SUM/SUBTOTAL de intervalos ou de
    colunas de Tabela e aritmética (+ - * /, ex. "=5.92+3.95" escrito à mão
    numa linha). O que não se sabe calcular (ex. #REF!) torna o resultado
    None - nunca um total errado."""

    def __init__(self, wb):
        self.wb = wb
        self._memo: Dict[tuple, float] = {}

    def valor(self, ws, coord: str) -> Optional[float]:
        try:
            return self._celula(ws, coord, 0)
        except (_Incalculavel, RecursionError, KeyError, ValueError, ZeroDivisionError, SyntaxError):
            return None

    def _celula(self, ws, coord: str, prof: int) -> float:
        chave = (ws.title, coord)
        if chave in self._memo:
            return self._memo[chave]
        v = ws[coord].value
        if v is None or (isinstance(v, str) and not v.startswith("=")):
            return 0.0  # vazio ou texto: o Excel ignora na soma
        if isinstance(v, bool):
            return float(v)
        if isinstance(v, (int, float)):
            return float(v)
        if not isinstance(v, str) or prof > 30:
            raise _Incalculavel(v)
        resultado = self._expressao(ws, v[1:], prof)
        self._memo[chave] = resultado
        return resultado

    def _expressao(self, ws, f: str, prof: int) -> float:
        f = f.strip()
        if "#" in f:  # #REF!, #N/A...
            raise _Incalculavel(f)
        while _FUNCAO.search(f):  # de dentro para fora, se houver funções dentro de funções
            f = _FUNCAO.sub(lambda m: repr(self._funcao(ws, m.group(1).upper(), m.group(2), prof)), f)
        f = _REF_FOLHA.sub(lambda m: repr(self._celula(self._folha(m.group(1) or m.group(2)),
                                                       f"{m.group(3)}{m.group(4)}", prof + 1)), f)
        f = _REF_LOCAL.sub(lambda m: repr(self._celula(ws, f"{m.group(1)}{m.group(2)}", prof + 1)), f)
        if not _ARITMETICA.match(f):
            raise _Incalculavel(f)
        return float(eval(f, {"__builtins__": {}}, {}))  # só números e + - * / ( )

    def _folha(self, nome: str):
        if nome not in self.wb.sheetnames:
            raise _Incalculavel(nome)
        return self.wb[nome]

    def _funcao(self, ws, nome: str, args: str, prof: int) -> float:
        partes = [a.strip() for a in args.split(",")]
        if nome == "SUBTOTAL":
            partes = partes[1:]  # o 1.º argumento é o tipo (9/109 = soma)
        total = 0.0
        for parte in partes:
            if re.fullmatch(r"-?[0-9.]+(?:e-?\d+)?", parte):  # já calculado (função dentro de função)
                total += float(parte)
                continue
            m = re.match(r"^([A-Za-z0-9_]+)\[([^\]]+)\]$", parte)
            if m:
                total += self._coluna_tabela(m.group(1), m.group(2), prof)
                continue
            folha = ws
            if "!" in parte:
                nome_folha, parte = parte.rsplit("!", 1)
                folha = self._folha(nome_folha.strip("'"))
            c1, r1, c2, r2 = range_boundaries(parte.replace("$", ""))
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    total += self._celula(folha, f"{get_column_letter(c)}{r}", prof + 1)
        return total

    def _coluna_tabela(self, nome: str, coluna: str, prof: int) -> float:
        dono = next((w for w in self.wb.worksheets if nome in w.tables), None)
        if dono is None:
            raise _Incalculavel(nome)
        tabela = dono.tables[nome]
        c1, r1, c2, r2 = range_boundaries(tabela.ref)
        # a fórmula pode chegar sem espaços ("Valor Pago" -> "ValorPago"): comparar sem eles
        sem_espacos = lambda texto: re.sub(r"\s+", "", str(texto or "")).lower()
        cabecalho = [sem_espacos(dono.cell(r1, c).value) for c in range(c1, c2 + 1)]
        if sem_espacos(coluna) not in cabecalho:
            raise _Incalculavel(coluna)
        col = get_column_letter(c1 + cabecalho.index(sem_espacos(coluna)))
        fim = r2 - (tabela.totalsRowCount or 0)
        return sum(self._celula(dono, f"{col}{r}", prof + 1) for r in range(r1 + 1, fim + 1))


def _linhas_de_texto(valor) -> int:
    return len([l for l in str(valor or "").splitlines() if l.strip()])


def _texto(valor) -> str:
    return re.sub(r"\s+", " ", str(valor or "")).strip().lower()


def colunas_do_resumo(ws) -> Optional[Dict[str, Any]]:
    """O Resumo mudou de forma ao longo do ano (jan-mai: Recebido/Total
    RECEBIDO, Pago/Valor PAGO e um só Saldo Bancário; jun: Escrituras e CPCV
    à cabeça; desde jul: CPCVs/escrituras com Valor Recebido e saldo
    disponível/indisponível) - as colunas encontram-se pelo nome no
    cabeçalho. Devolve {campo: letra} e a linha do cabeçalho, ou None."""
    from openpyxl.utils import get_column_letter

    linha = next((r for r in range(1, 7) if any(_texto(ws.cell(r, c).value) == "previsão" for c in range(2, 20))), None)
    if linha is None:
        return None
    cab = {c: _texto(ws.cell(linha, c).value) for c in range(2, 25)}
    grupo = {c: _texto(ws.cell(linha - 1, c).value) for c in range(2, 25)} if linha > 1 else {}
    inicio_pag = next((c for c, g in grupo.items() if g.startswith("pagamentos")), None)
    if inicio_pag is None:
        return None
    letra = get_column_letter
    achar = lambda pred, de=2, ate=25: next((c for c in range(de, ate) if pred(cab.get(c, ""), grupo.get(c, ""))), None)

    rec = achar(lambda h, g: h == "total recebido", ate=inicio_pag) or achar(lambda h, g: h == "recebido", ate=inicio_pag)
    pag = achar(lambda h, g: h == "valor pago", de=inicio_pag) or achar(lambda h, g: h == "pago", de=inicio_pag)
    if not rec or not pag:
        return None
    colunas: Dict[str, Any] = {"linha": linha, "recebido": letra(rec), "pago": letra(pag)}
    disp = achar(lambda h, g: ("dispon" in h or "dispon" in g) and "indispon" not in h + g)
    indisp = achar(lambda h, g: "indispon" in h or "indispon" in g)
    saldo = achar(lambda h, g: "bancário" in h + " " + g and "dispon" not in h + g, de=inicio_pag)
    for campo, c in (("saldo_disponivel", disp), ("saldo_indisponivel", indisp), ("saldo", saldo)):
        if c:
            colunas[campo] = letra(c)
    for tipo, nomes in (("cpcv", ("# de cpcvs", "cpcv")), ("escritura", ("# de escrituras", "escrituras"))):
        lista = achar(lambda h, g, n=nomes[0]: h == n, ate=inicio_pag)
        if lista:  # desde jul: a lista e, à frente, o Valor Recebido
            valor = achar(lambda h, g: h == "valor recebido", de=lista + 1, ate=inicio_pag)
            colunas[f"n_{tipo}"] = letra(lista)
            if valor:
                colunas[tipo] = letra(valor)
        else:  # junho: a coluna "CPCV"/"Escrituras" é o número deles, não o valor
            c = achar(lambda h, g, n=nomes[1]: h == n, ate=inicio_pag)
            if c:
                colunas[f"contagem_{tipo}"] = letra(c)
    return colunas


def _vendas(ws, calc, colunas: Dict[str, Any], tipo: str, r: int) -> Dict[str, Any]:
    """{tipo: valor recebido ou None, n_tipo: quantos ou None} - None quando
    o Resumo desse mês não tem essa informação."""
    valor = calc.valor(ws, f"{colunas[tipo]}{r}") if tipo in colunas else None
    if f"n_{tipo}" in colunas:
        n = _linhas_de_texto(ws[f"{colunas[f'n_{tipo}']}{r}"].value)
    elif f"contagem_{tipo}" in colunas:
        n = int(calc.valor(ws, f"{colunas[f'contagem_{tipo}']}{r}") or 0)
    else:
        n = None
    return {tipo: valor, f"n_{tipo}": n}


def ler_resumo(caminho: str) -> Optional[Dict[int, Dict[str, Any]]]:
    """{dia: {recebido, pago, saldo_disponivel, saldo_indisponivel,
    saldo_bancario, cpcv, n_cpcv, escritura, n_escritura}} do Resumo deste
    ficheiro (None num valor que o Resumo não tem ou que dá erro), ou None se
    não existe/não se reconhece o formato. Guardado em memória enquanto o
    ficheiro não mudar."""
    try:
        estado = os.stat(caminho)
    except OSError:
        return None
    chave = (caminho, estado.st_mtime_ns, estado.st_size)
    if chave not in _CACHE:
        if len(_CACHE) > 24:
            _CACHE.clear()
        _CACHE[chave] = _ler(caminho)
    return _CACHE[chave]


def _ler(caminho: str) -> Optional[Dict[int, Dict[str, Any]]]:
    try:
        wb = openpyxl.load_workbook(caminho)
    except Exception:
        logger.warning("Não consegui abrir o Mapa %s", caminho, exc_info=True)
        return None
    colunas = colunas_do_resumo(wb["Resumo"]) if "Resumo" in wb.sheetnames else None
    if not colunas:
        logger.warning("Resumo do Mapa %s sem o cabeçalho esperado - não o leio", caminho)
        return None
    ws, calc = wb["Resumo"], _Calculadora(wb)

    def ler(campo: str, r: int) -> Optional[float]:
        coluna = colunas.get(campo)
        if not coluna or ws[f"{coluna}{r}"].value is None and campo.startswith("saldo"):
            return None  # sem esse valor no Resumo, ou sem saldo nessa linha (fim de semana)
        return calc.valor(ws, f"{coluna}{r}")

    dias: Dict[int, Dict[str, Any]] = {}
    for r in range(colunas["linha"] + 1, ws.max_row + 1):
        dia = ws[f"A{r}"].value
        if not isinstance(dia, int) or not 1 <= dia <= 31:
            continue
        disponivel, indisponivel, saldo = ler("saldo_disponivel", r), ler("saldo_indisponivel", r), ler("saldo", r)
        dias[dia] = {
            "recebido": ler("recebido", r), "pago": ler("pago", r),
            "saldo_disponivel": disponivel, "saldo_indisponivel": indisponivel,
            "saldo_bancario": (disponivel + (indisponivel or 0.0)) if disponivel is not None else saldo,
            **{chave: valor for tipo in ("cpcv", "escritura") for chave, valor in _vendas(ws, calc, colunas, tipo, r).items()},
        }
    return dias


def resumo_do_mes(ano: int, mes: int) -> Optional[Dict[date, Dict[str, Any]]]:
    """O Resumo do Mapa desse mês por data, ou None (sem Mapa/sem acesso)."""
    try:
        caminho = caminho_mapa(date(ano, mes, 1))
    except Exception:  # ONEDRIVE_RAIZ em falta, ...
        return None
    dias = ler_resumo(caminho)
    if not dias:
        return None
    resultado = {}
    for d, v in dias.items():
        try:
            resultado[date(ano, mes, d)] = dict(v)  # cópia: quem chama pode mexer, a cache não
        except ValueError:  # 31 num mês de 30 dias
            continue
    return resultado
